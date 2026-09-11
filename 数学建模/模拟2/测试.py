import pandas as pd
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Inches
import os

# 文件路径
forecast_csv = r"D:\PythonProject1\数学建模\模拟2\资料\bike_forecast.csv"
output_docx = r"D:\PythonProject1\数学建模\模拟2\第一题预测论文自动生成.docx"
img_folder = r"D:\PythonProject1\数学建模\模拟2\figures"

os.makedirs(img_folder, exist_ok=True)

# 读取预测数据
forecast = pd.read_csv(forecast_csv)

# 创建 Word 文档
doc = Document()
doc.add_heading('校园共享单车调度与维护问题 — 第一题', 0)
doc.add_paragraph('姓名：XXX')
doc.add_paragraph('学号：XXX')
doc.add_paragraph('指导老师：XXX')
doc.add_paragraph('学校：XXX')
doc.add_paragraph('日期：2026年5月6日')
doc.add_page_break()

# 摘要
doc.add_heading('摘要', level=1)
doc.add_paragraph(
    "本文针对校园共享单车的高峰需求预测问题，基于学校作息时间和历史单车使用数据，"
    "对各站点在早、中、晚三个高峰期的单车需求进行了预测。结果显示，教学楼、食堂和主要校门为高需求区域，为后续调度优化提供数据依据。"
)

# 研究背景与数据预处理
doc.add_heading('研究背景与问题描述', level=1)
doc.add_paragraph(
    "校园共享单车广泛应用于高校，方便学生在校内快速通行。然而，高峰期单车供应不足会导致用户等待和使用体验下降。"
    "为了科学调度车辆，需要提前预测每个站点在不同时间段的单车需求量。"
)
doc.add_heading('数据来源与预处理', level=2)
doc.add_paragraph(
    "- 数据来源：附件1-共享单车分布统计表，附件3-学校作息时间表\n"
    "- 数据内容：站点名称、时间、车辆数\n"
    "- 预处理：非数值处理，自动识别时间列，高峰时间段分类"
)

# 方法
doc.add_heading('方法', level=1)
doc.add_paragraph(
    "将每个时间点分配到对应时间段，并对每个站点每个时间段的车辆数求平均作为预测需求。"
)

# 定义高峰时间段
peak_periods = {
    'morning_peak': ["07:30:00", "08:50:00", "09:50:00"],
    'lunch_peak': ["12:20:00", "13:50:00"],
    'evening_peak': ["18:00:00", "19:00:00"]
}

# 绘制柱状图并保存图片
for period, times in peak_periods.items():
    df_period = forecast[forecast['time_period'] == period].copy()
    df_period = df_period.sort_values(by='predicted_demand', ascending=False)

    plt.figure(figsize=(14, 6))
    bars = plt.bar(df_period['station_id'], df_period['predicted_demand'], color='skyblue')
    plt.xticks(rotation=45, ha='right')
    plt.ylabel("Predicted Bikes")
    plt.title(f"{period} Demand Forecast per Station")
    for bar in bars:
        plt.text(bar.get_x() + bar.get_width() / 2, bar.get_height() + 1, f'{int(bar.get_height())}', ha='center',
                 va='bottom', fontsize=9)
    plt.tight_layout()

    img_path = os.path.join(img_folder, f"{period}.png")
    plt.savefig(img_path)
    plt.close()

    # 插入图片到 Word
    doc.add_heading(f'{period}预测需求柱状图', level=2)
    doc.add_picture(img_path, width=Inches(6.5))

# 结果表格占位
doc.add_heading('预测结果示例表格', level=1)
table = doc.add_table(rows=1, cols=3)
hdr_cells = table.rows[0].cells
hdr_cells[0].text = 'station_id'
hdr_cells[1].text = 'time_period'
hdr_cells[2].text = 'predicted_demand'
row_cells = table.add_row().cells
row_cells[0].text = '教学楼'
row_cells[1].text = 'morning_peak'
row_cells[2].text = '45'

# 讨论与总结
doc.add_heading('讨论', level=1)
doc.add_paragraph(
    "- 方法优势：简单高效，快速生成预测\n"
    "- 局限性：时间段粗，不能反映分钟级变化，对突发事件无预测\n"
    "- 改进方向：使用回归或时间序列模型，结合天气或课程安排动态预测"
)
doc.add_heading('总结', level=1)
doc.add_paragraph(
    "本文完成了校园共享单车高峰期需求预测，生成预测表格和可视化柱状图，为第二题调度优化提供数据支持。"
)

# 保存 Word 文档
doc.save(output_docx)
print("Word 文档已生成:", output_docx)